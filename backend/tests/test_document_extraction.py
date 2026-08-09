"""Deterministic parser, normalization, persistence, and API tests."""

import asyncio
import logging
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select

import app.services.knowledge_document as knowledge_service_module
from app.core.config import Settings
from app.document_parsers.docx import DocxDocumentParser
from app.document_parsers.exceptions import (
    CorruptedDocumentError,
    EncryptedPdfError,
    ExtractionLimitExceededError,
    NoExtractableTextError,
)
from app.document_parsers.normalization import normalize_text
from app.document_parsers.pdf import PdfDocumentParser
from app.document_parsers.registry import DocumentParserRegistry
from app.document_parsers.text import MarkdownDocumentParser, TextDocumentParser
from app.models.knowledge_document import KnowledgeDocument
from app.models.knowledge_document_content import KnowledgeDocumentContent
from app.services.knowledge_document import KnowledgeDocumentService
from app.storage.local import LocalStorageProvider
from tests.conftest import (
    FakeEmailProvider,
    TestSessionLocal,
    register_verify_and_login,
)

BASE_URL = "/api/v1/knowledge"


def _auth(client: TestClient, email: str) -> dict[str, str]:
    _, token = register_verify_and_login(client, email=email)
    return {"Authorization": f"Bearer {token}"}


def _pdf(*pages: str, encrypted: bool = False) -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        stream = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("password")
    writer.write(output)
    return output.getvalue()


def _docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "A2"
    table.cell(1, 0).text = "B1"
    table.cell(1, 1).text = "B2"
    document.add_paragraph("Last paragraph")
    document.save(output)
    return output.getvalue()


def test_text_markdown_bom_and_thai_extraction() -> None:
    text = TextDocumentParser().extract("สวัสดี\r\nOneInFlow".encode())
    bom = TextDocumentParser().extract(b"\xef\xbb\xbfBOM text")
    markdown = MarkdownDocumentParser().extract(b"# Heading\n\n- item")
    assert "สวัสดี" in text
    assert bom == "BOM text"
    assert markdown == "# Heading\n\n- item"


def test_docx_paragraph_and_flattened_table_order() -> None:
    text = DocxDocumentParser().extract(_docx())
    assert text.splitlines() == [
        "First paragraph",
        "A1\tA2",
        "B1\tB2",
        "Last paragraph",
    ]


def test_pdf_text_and_page_order() -> None:
    text = PdfDocumentParser().extract(_pdf("Page one", "Page two"))
    assert text.index("Page one") < text.index("Page two")


def test_encrypted_corrupted_and_empty_documents_are_typed_failures() -> None:
    with pytest.raises(EncryptedPdfError):
        PdfDocumentParser().extract(_pdf("Secret", encrypted=True))
    with pytest.raises(CorruptedDocumentError):
        PdfDocumentParser().extract(b"%PDF-corrupted")
    with pytest.raises(CorruptedDocumentError):
        DocxDocumentParser().extract(b"PK\x03\x04corrupted")
    with pytest.raises(NoExtractableTextError):
        DocumentParserRegistry().extract(
            mime_type="application/pdf",
            data=_pdf(""),
            maximum_characters=100,
        )


def test_normalization_rules_and_character_limit() -> None:
    normalized = normalize_text(
        "ไทย\r\nline  \r\x00\n\n\n\nend", maximum_characters=100
    )
    assert normalized == "ไทย\nline\n\n\nend"
    with pytest.raises(ExtractionLimitExceededError):
        normalize_text("12345", maximum_characters=4)


@pytest.fixture
def extraction_storage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> Path:
    monkeypatch.setattr(
        knowledge_service_module,
        "create_storage_provider",
        lambda settings: LocalStorageProvider(tmp_path),
    )
    return tmp_path


def test_upload_persists_one_content_row_and_returns_counts(
    client_db: TestClient,
    fake_email_provider: FakeEmailProvider,
    extraction_storage: Path,
) -> None:
    headers = _auth(client_db, "extraction@example.com")
    response = client_db.post(
        f"{BASE_URL}/upload",
        headers=headers,
        files={"file": ("guide.txt", b"one two\nthree", "text/plain")},
    )
    assert response.status_code == 201
    document = response.json()["data"]
    assert document["status"] == "READY"
    content_response = client_db.get(
        f"{BASE_URL}/{document['id']}/content", headers=headers
    )
    assert content_response.status_code == 200
    content = content_response.json()["data"]
    assert content["extracted_text"] == "one two\nthree"
    assert content["character_count"] == len("one two\nthree")
    assert content["word_count"] == 3
    assert "storage_key" not in content
    assert "path" not in content


def test_retry_upserts_without_duplicate_content(
    client_db: TestClient,
    fake_email_provider: FakeEmailProvider,
    extraction_storage: Path,
) -> None:
    headers = _auth(client_db, "retry@example.com")
    uploaded = client_db.post(
        f"{BASE_URL}/upload",
        headers=headers,
        files={"file": ("retry.txt", b"retry content", "text/plain")},
    ).json()["data"]

    async def _retry() -> tuple[int, str]:
        async with TestSessionLocal() as session:
            document = await session.get(KnowledgeDocument, uploaded["id"])
            assert document is not None
            service = KnowledgeDocumentService(
                session, storage=LocalStorageProvider(extraction_storage)
            )
            await service.extract(document)
            count = await session.scalar(
                select(func.count(KnowledgeDocumentContent.id)).where(
                    KnowledgeDocumentContent.document_id == document.id
                )
            )
            return int(count or 0), document.status.value

    count, status = asyncio.run(_retry())
    assert count == 1
    assert status == "READY"


def test_failed_extraction_sets_failed_without_content_leak(
    client_db: TestClient,
    fake_email_provider: FakeEmailProvider,
    extraction_storage: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    headers = _auth(client_db, "failed@example.com")
    secret_content = "private-document-content"
    settings = Settings(
        _env_file=None,
        KNOWLEDGE_MAX_EXTRACTED_CHARACTERS=4,
    )
    original_get_settings = knowledge_service_module.get_settings
    knowledge_service_module.get_settings = lambda: settings
    try:
        with caplog.at_level(logging.WARNING):
            response = client_db.post(
                f"{BASE_URL}/upload",
                headers=headers,
                files={
                    "file": ("failed.txt", secret_content.encode(), "text/plain")
                },
            )
    finally:
        knowledge_service_module.get_settings = original_get_settings
    assert response.status_code == 422
    listing = client_db.get(BASE_URL, headers=headers).json()["data"]["items"]
    assert listing[0]["status"] == "FAILED"
    assert secret_content not in caplog.text


def test_content_endpoint_enforces_workspace_isolation(
    client_db: TestClient,
    fake_email_provider: FakeEmailProvider,
    extraction_storage: Path,
) -> None:
    first = _auth(client_db, "content-owner@example.com")
    document = client_db.post(
        f"{BASE_URL}/upload",
        headers=first,
        files={"file": ("private.txt", b"private", "text/plain")},
    ).json()["data"]
    second = _auth(client_db, "content-other@example.com")
    response = client_db.get(f"{BASE_URL}/{document['id']}/content", headers=second)
    assert response.status_code == 404
